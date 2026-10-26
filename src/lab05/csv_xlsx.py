"""CSV → XLSX через openpyxl."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Первая строка — заголовок, лист "Sheet1", автоширина колонок (≥ 8).
    """
    src = Path(csv_path)
    if not src.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    with src.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))

    if not rows or not any(rows):
        raise ValueError("Пустой CSV")
    if not rows[0]:
        raise ValueError("CSV без заголовка")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    widths = [0] * len(rows[0])
    for i, row in enumerate(rows):
        # выравниваем длину строк по количеству колонок заголовка
        row = list(row) + [""] * (len(widths) - len(row))
        ws.append(row[: len(widths)])
        for j, value in enumerate(row[: len(widths)]):
            widths[j] = max(widths[j], len(str(value)))

    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(w + 2, 8)

    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
